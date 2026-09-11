from __future__ import annotations

import hashlib
import json
from pathlib import Path
import socket

from docx import Document
import pytest
from pypdf import PdfReader

from services import job as job_mod
from services.meeting_archive import (
    archive_meeting,
    rename_meeting_archive,
    sanitize_windows_component,
)
from tests.helpers import generate_sine_wav


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_job(tmp_path: Path) -> tuple[Path, dict[str, str]]:
    job_dir = job_mod.create_job("archive-fixture", str(tmp_path / "jobs"))
    source_path = job_dir / "source" / "授权录音.wav"
    generate_sine_wav(str(source_path), duration_sec=1.0, framerate=8000)

    provider_hashes = {}
    provider_outputs = {}
    for provider_code in ("M3ASR", "gptTr", "Gem35T"):
        evidence_dir = job_dir / "quality" / "provider-evidence" / provider_code
        evidence_dir.mkdir(parents=True)
        json_path = evidence_dir / f"{provider_code}.json"
        html_path = evidence_dir / f"{provider_code}.html"
        json_path.write_text(
            json.dumps(
                {"provider_code": provider_code, "text": f"{provider_code} 原始证据"},
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        html_path.write_text(f"<html>{provider_code} 原始证据</html>", encoding="utf-8")
        provider_hashes[f"{provider_code}.json"] = _sha256(json_path)
        provider_hashes[f"{provider_code}.html"] = _sha256(html_path)
        provider_outputs[provider_code] = {"json": str(json_path), "html": str(html_path)}

    final_reviewed = {
        "schema_version": "apma.final-reviewed.v1",
        "statistics": {
            "total_windows": 2,
            "resolved_windows": 2,
            "unresolved_windows": 0,
        },
        "windows": [
            {
                "window_id": "window-1",
                "global_start_sec": 915.25,
                "global_end_sec": 935.5,
                "final_text": "最终中文文本。Hello 世界。",
                "review_required": False,
                "provider_candidates": {"M3ASR": {"text": "候选一"}},
                "review_resolution": {
                    "decision_type": "provider_candidate",
                    "selected_provider": "M3ASR",
                },
            },
            {
                "window_id": "window-2",
                "global_start_sec": 935.5,
                "global_end_sec": 960.0,
                "final_text": "Selamat pagi, lí hó; meeting selesai.",
                "review_required": False,
                "provider_candidates": {
                    "gptTr": {"text": "Selamat pagi, lí hó; meeting selesai."}
                },
                "review_resolution": {
                    "decision_type": "manual_correction",
                    "selected_provider": None,
                },
            },
        ],
    }
    final_path = job_dir / "quality" / "review" / "FINAL_REVIEWED.json"
    final_path.parent.mkdir(parents=True)
    final_path.write_text(json.dumps(final_reviewed, indent=2, ensure_ascii=False), encoding="utf-8")

    overlay = {
        "schema_version": "apma.speaker-overlay.v1",
        "windows": [
            {
                "window_id": "window-1",
                "speaker_evidence": {
                    "segments": [
                        {
                            "local_speaker_key": "local-a",
                            "speaker_id": "spk:0",
                            "global_start_sec": 915.25,
                            "global_end_sec": 935.5,
                        }
                    ]
                },
            },
            {
                "window_id": "window-2",
                "speaker_evidence": {
                    "segments": [
                        {
                            "local_speaker_key": "local-b",
                            "speaker_id": "spk:1",
                            "global_start_sec": 935.5,
                            "global_end_sec": 960.0,
                        }
                    ]
                },
            },
        ],
    }
    overlay_path = job_dir / "quality" / "speaker-overlay" / "speaker_overlay.json"
    overlay_path.parent.mkdir(parents=True)
    overlay_path.write_text(json.dumps(overlay, indent=2), encoding="utf-8")
    mappings = {
        "schema_version": "apma.speaker-mappings.v1",
        "mappings": [
            {
                "local_speaker_key": "local-a",
                "provider_speaker_id": "spk:0",
                "canonical_speaker_id": "canonical-speaker-001",
                "display_name": "John",
                "mapping_source": "human",
            }
        ],
    }
    mapping_path = job_dir / "quality" / "review" / "speaker_mappings.json"
    mapping_path.write_text(json.dumps(mappings, indent=2), encoding="utf-8")

    manifest = job_mod.read_manifest(job_dir)
    manifest["source"] = {
        "path": str(source_path),
        "filename": source_path.name,
        "sha256": _sha256(source_path),
    }
    manifest["quality"] = {
        "outputs": {
            "providers": provider_outputs,
            "final_reviewed": {"json": str(final_path)},
            "speaker_overlay": {
                "json": str(overlay_path),
                "speaker_mappings_json": str(mapping_path),
            },
        }
    }
    job_mod.write_manifest(job_dir, manifest)
    return job_dir, provider_hashes


def _archive(job_dir: Path, archive_root: Path, **overrides):
    values = {
        "meeting_id": "meeting-stable-001",
        "meeting_start": "2026-07-12T10:00:00+08:00",
        "meeting_datetime_source": "user",
        "likely_project": "Orchid Research",
        "content": "Quarterly Review",
        "short_name": "Orchid Review",
    }
    values.update(overrides)
    return archive_meeting(job_dir, archive_root, **values)


def test_archive_contract_provider_names_and_final_exports(tmp_path, monkeypatch):
    job_dir, provider_hashes = _write_job(tmp_path)
    archive_root = tmp_path / "archive"

    def fail_network(*_args, **_kwargs):
        raise AssertionError("Meeting archive attempted network access")

    monkeypatch.setattr(socket, "create_connection", fail_network)
    result = _archive(job_dir, archive_root)
    folder = Path(result["archive_folder"])

    assert folder.name == "2607121000 12Jul2026 Orchid Research Quarterly Review"
    assert {path.name for path in folder.iterdir()} == {"source", "asr", "final", "manifest.json"}
    assert result["provider_filenames"] == {
        provider: {
            "json": f"2607121000 Orchid Review {provider}.json",
            "html": f"2607121000 Orchid Review {provider}.html",
        }
        for provider in ("M3ASR", "gptTr", "Gem35T")
    }
    assert result["final_filenames"] == {
        extension: f"2607121000 Orchid Review FINAL.{extension}"
        for extension in ("json", "html", "txt", "srt", "vtt", "docx", "pdf")
    }
    assert result["formats_completed"] == [
        "JSON",
        "HTML",
        "TXT",
        "SRT",
        "VTT",
        "DOCX",
        "PDF",
    ]
    assert "canonical FINAL JSON" in result["docx_status"]
    assert "canonical FINAL JSON" in result["pdf_status"]

    for source_name, expected_hash in provider_hashes.items():
        source = next(job_dir.rglob(source_name))
        assert _sha256(source) == expected_hash
        archived = next((folder / "asr").glob(f"* {source_name}"))
        assert _sha256(archived) == expected_hash

    final_json = folder / "final" / result["final_filenames"]["json"]
    canonical = json.loads(final_json.read_text(encoding="utf-8"))
    txt = (folder / "final" / result["final_filenames"]["txt"]).read_text(encoding="utf-8")
    html = (folder / "final" / result["final_filenames"]["html"]).read_text(encoding="utf-8")
    srt = (folder / "final" / result["final_filenames"]["srt"]).read_text(encoding="utf-8")
    vtt = (folder / "final" / result["final_filenames"]["vtt"]).read_text(encoding="utf-8")
    docx_path = folder / "final" / result["final_filenames"]["docx"]
    pdf_path = folder / "final" / result["final_filenames"]["pdf"]
    document = Document(docx_path)
    docx_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(pdf_path).pages)
    assert txt == canonical["text"] + "\n"
    assert all(segment["text"] in html for segment in canonical["segments"])
    assert "00:15:15,250 --> 00:15:35,500" in srt
    assert "00:15:15.250 --> 00:15:35.500" in vtt
    assert "John [spk:0]" in canonical["segments"][0]["speaker"]
    assert "John [spk:0]" in html
    assert "John [spk:0]" in srt
    assert canonical["segments"][1]["speaker"] == "spk:1"
    for expected in ("最终中文文本。Hello 世界。", "Selamat pagi", "lí hó"):
        assert expected in txt
        assert expected in docx_text
        assert expected in pdf_text
    for expected in ("John [spk:0]", "00:15:15.250 --> 00:15:35.500"):
        assert expected in docx_text
        assert expected in pdf_text
    assert canonical["authority"]["canonical"] is True
    assert canonical["authority"]["derived_exports"] == [
        "HTML",
        "TXT",
        "SRT",
        "VTT",
        "DOCX",
        "PDF",
    ]
    assert result["final_filenames"]["json"] in html


def test_rerun_creates_non_destructive_revision(tmp_path):
    job_dir, _ = _write_job(tmp_path)
    archive_root = tmp_path / "archive"
    first = _archive(job_dir, archive_root)
    first_folder = Path(first["archive_folder"])
    first_files = {
        path.relative_to(first_folder): _sha256(path)
        for path in first_folder.rglob("*")
        if path.is_file() and path.name != "manifest.json"
    }

    second = _archive(job_dir, archive_root)
    assert second["revision"] == 2
    assert second["archive_folder"] == first["archive_folder"]
    assert second["provider_filenames"]["M3ASR"]["json"] == (
        "2607121000 Orchid Review rev02 M3ASR.json"
    )
    assert second["final_filenames"]["json"] == "2607121000 Orchid Review rev02 FINAL.json"
    for relative, original_hash in first_files.items():
        assert _sha256(first_folder / relative) == original_hash
    manifest = json.loads((first_folder / "manifest.json").read_text(encoding="utf-8"))
    assert [item["revision"] for item in manifest["revisions"]] == [1, 2]


def test_stable_meeting_id_survives_human_readable_rename(tmp_path):
    job_dir, _ = _write_job(tmp_path)
    archive_root = tmp_path / "archive"
    first = _archive(job_dir, archive_root)
    old_folder = Path(first["archive_folder"])
    renamed = rename_meeting_archive(
        archive_root,
        meeting_id="meeting-stable-001",
        meeting_start="2026-07-12T10:00:00+08:00",
        meeting_datetime_source="user",
        likely_project="研究/项目",
        content="估值:会议?",
        short_name="研究*估值",
    )
    assert not old_folder.exists()
    assert renamed.name == "2607121000 12Jul2026 研究 项目 估值 会议"
    manifest = json.loads((renamed / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["meeting_id"] == "meeting-stable-001"
    assert manifest["meeting_metadata"]["short_name"] == "研究 估值"
    for revision in manifest["revisions"]:
        for values in revision["provider_outputs"].values():
            for item in values.values():
                assert (renamed / item["relative_path"]).is_file()
        for item in revision["final_outputs"].values():
            assert (renamed / item["relative_path"]).is_file()


def test_folder_collision_and_unknown_time_do_not_overwrite_or_invent(tmp_path):
    job_dir, _ = _write_job(tmp_path)
    archive_root = tmp_path / "archive"
    collision = archive_root / "2607121000 12Jul2026 Orchid Research Quarterly Review"
    collision.mkdir(parents=True)
    (collision / "manifest.json").write_text(
        json.dumps({"meeting_id": "different-meeting"}), encoding="utf-8"
    )
    result = _archive(job_dir, archive_root)
    assert Path(result["archive_folder"]).name.endswith(" (2)")
    assert json.loads((collision / "manifest.json").read_text(encoding="utf-8")) == {
        "meeting_id": "different-meeting"
    }

    other_root = tmp_path / "unknown-archive"
    unknown = _archive(
        job_dir,
        other_root,
        meeting_id="meeting-unknown-time",
        meeting_start=None,
        meeting_datetime_source="unknown",
    )
    assert Path(unknown["archive_folder"]).name.startswith("UnknownDateTime UnknownDate ")
    manifest = json.loads(
        (Path(unknown["archive_folder"]) / "manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["meeting_metadata"]["meeting_start"] is None
    assert manifest["meeting_metadata"]["meeting_datetime_source"] == "unknown"


def test_windows_sanitizer_preserves_unicode_and_removes_illegal_endings():
    value = sanitize_windows_component('  中文/项目:<会议>?* .  ')
    assert value == "中文 项目 会议"
    assert not any(character in value for character in '<>:"/\\|?*')
    assert not value.endswith((" ", "."))
    assert sanitize_windows_component("CON") == "_CON"


def test_date_only_metadata_cannot_silently_invent_midnight(tmp_path):
    job_dir, _ = _write_job(tmp_path)
    with pytest.raises(ValueError, match="explicit hour and minute"):
        _archive(job_dir, tmp_path / "archive", meeting_start="2026-07-12")
