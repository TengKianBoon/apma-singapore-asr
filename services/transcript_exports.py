from __future__ import annotations

import math
from html import escape
from pathlib import Path
from typing import Any, Iterable


EXPORT_FILENAMES = {
    "full_transcript_html": "full_transcript.html",
    "full_transcript_srt": "full_transcript.srt",
    "full_transcript_vtt": "full_transcript.vtt",
}


def _write_text(path: Path, content: str) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(content)


def _as_finite_float(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _timed_segments(payload: dict[str, Any]) -> Iterable[tuple[int, int, dict[str, Any]]]:
    segments = payload.get("segments")
    if not isinstance(segments, list):
        return

    for segment in segments:
        if not isinstance(segment, dict):
            continue
        text = str(segment.get("text", "")).strip()
        start = _as_finite_float(segment.get("start_sec"))
        end = _as_finite_float(segment.get("end_sec"))
        if not text or start is None or end is None or start < 0 or end <= start:
            continue
        start_ms = int(round(start * 1000))
        end_ms = int(round(end * 1000))
        if end_ms <= start_ms:
            continue
        yield start_ms, end_ms, segment


def _format_timestamp(milliseconds: int, decimal_separator: str) -> str:
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    seconds, millis = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}{decimal_separator}{millis:03d}"


def _caption_text(segment: dict[str, Any]) -> str:
    text = str(segment.get("text", "")).strip()
    speaker = str(segment.get("speaker") or "").strip()
    rendered = f"{speaker}: {text}" if speaker else text
    return escape(rendered, quote=False)


def _segment_timing(segment: dict[str, Any]) -> str:
    start = _as_finite_float(segment.get("start_sec"))
    end = _as_finite_float(segment.get("end_sec"))
    if start is None or end is None or start < 0 or end <= start:
        return "Timing unavailable"
    return (
        f"{_format_timestamp(int(round(start * 1000)), '.')}"
        f" --> {_format_timestamp(int(round(end * 1000)), '.')}"
    )


def render_srt(payload: dict[str, Any]) -> str:
    cues = []
    for cue_number, (start_ms, end_ms, segment) in enumerate(_timed_segments(payload), start=1):
        cues.append(
            f"{cue_number}\n"
            f"{_format_timestamp(start_ms, ',')} --> {_format_timestamp(end_ms, ',')}\n"
            f"{_caption_text(segment)}"
        )
    return "\n\n".join(cues) + ("\n" if cues else "")


def render_vtt(payload: dict[str, Any]) -> str:
    cues = ["WEBVTT"]
    for cue_number, (start_ms, end_ms, segment) in enumerate(_timed_segments(payload), start=1):
        cues.append(
            f"{cue_number}\n"
            f"{_format_timestamp(start_ms, '.')} --> {_format_timestamp(end_ms, '.')}\n"
            f"{_caption_text(segment)}"
        )
    return "\n\n".join(cues) + "\n"


def _metadata_items(payload: dict[str, Any]) -> list[tuple[str, str]]:
    review = payload.get("review") if isinstance(payload.get("review"), dict) else {}
    provenance = payload.get("provenance") if isinstance(payload.get("provenance"), dict) else {}
    providers = provenance.get("providers") if isinstance(provenance.get("providers"), list) else []
    models = provenance.get("models") if isinstance(provenance.get("models"), list) else []
    source_hashes = (
        provenance.get("source_sha256")
        if isinstance(provenance.get("source_sha256"), list)
        else []
    )
    return [
        ("Job", str(payload.get("job_id") or "unknown")),
        ("Created", str(payload.get("created_at") or "unknown")),
        ("Review", str(review.get("highest_attention_grade") or "not graded").upper()),
        ("Providers", ", ".join(str(item) for item in providers) or "unknown"),
        ("Models", ", ".join(str(item) for item in models) or "unknown"),
        ("Source SHA-256", ", ".join(str(item) for item in source_hashes) or "unknown"),
        ("Segments", str(len(payload.get("segments") or []))),
    ]


def render_html(payload: dict[str, Any]) -> str:
    metadata = "\n".join(
        f'<div><dt>{escape(label)}</dt><dd>{escape(value)}</dd></div>'
        for label, value in _metadata_items(payload)
    )
    segment_rows = []
    for segment in payload.get("segments") or []:
        if not isinstance(segment, dict):
            continue
        text = str(segment.get("text", "")).strip()
        if not text:
            continue
        timing = _segment_timing(segment)
        speaker = str(segment.get("speaker") or "Unlabelled speaker")
        provider = str(segment.get("provider") or "unknown provider")
        model = str(segment.get("model") or "unknown model")
        segment_rows.append(
            '<article class="segment">'
            f'<div class="segment-meta"><strong>{escape(speaker)}</strong>'
            f'<span>{escape(timing)}</span></div>'
            f'<p>{escape(text)}</p>'
            f'<small>{escape(provider)} / {escape(model)}</small>'
            "</article>"
        )
    transcript = "\n".join(segment_rows) or '<p class="empty">No transcript segments available.</p>'
    job_id = escape(str(payload.get("job_id") or "unknown"))
    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; img-src 'none'; script-src 'none'; connect-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'">
  <title>APMA Transcript - {job_id}</title>
  <style>
    :root {{ color-scheme: light; font-family: Arial, sans-serif; color: #17202a; background: #f5f7f9; }}
    body {{ margin: 0; }}
    header, main {{ width: min(920px, calc(100% - 32px)); margin: 0 auto; }}
    header {{ padding: 32px 0 20px; }}
    h1 {{ margin: 0 0 8px; font-size: 28px; }}
    .authority {{ padding: 12px; border-left: 4px solid #1d6fdc; background: #eaf2fd; }}
    dl {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(180px, 1fr)); gap: 10px; }}
    dl div {{ border: 1px solid #d8dee6; background: white; padding: 10px; }}
    dt {{ color: #5d6875; font-size: 12px; }}
    dd {{ margin: 4px 0 0; overflow-wrap: anywhere; }}
    .segment {{ margin: 12px 0; padding: 14px; border: 1px solid #d8dee6; background: white; }}
    .segment-meta {{ display: flex; justify-content: space-between; gap: 12px; color: #3e4a59; }}
    .segment p {{ margin: 12px 0; white-space: pre-wrap; line-height: 1.5; }}
    small {{ color: #687483; }}
    .empty {{ background: white; padding: 16px; }}
  </style>
</head>
<body>
  <header>
    <h1>Meeting transcript</h1>
    <p class="authority"><strong>Authority notice:</strong> This HTML file is a derived human view. <code>full_transcript.json</code> is the authoritative transcript record.</p>
    <dl>{metadata}</dl>
  </header>
  <main>{transcript}</main>
</body>
</html>
"""


def write_docx(payload: dict[str, Any], path: Path) -> None:
    """Write a derived Unicode DOCX view from the canonical transcript payload."""

    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"Transcript export already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    document = Document()
    document.core_properties.title = "APMA Meeting Transcript"
    document.core_properties.subject = "Derived from APMA canonical FINAL JSON"
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    document.add_heading("Meeting transcript", level=0)
    authority = document.add_paragraph()
    authority.add_run("Authority notice: ").bold = True
    authority.add_run("This DOCX is a derived view. The corresponding FINAL.json is authoritative.")
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in _metadata_items(payload):
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    for segment in payload.get("segments") or []:
        if not isinstance(segment, dict):
            continue
        text = str(segment.get("text", "")).strip()
        if not text:
            continue
        heading = document.add_paragraph()
        speaker = str(segment.get("speaker") or "Unlabelled speaker")
        heading.add_run(f"{speaker} | {_segment_timing(segment)}").bold = True
        document.add_paragraph(text)

    try:
        document.save(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise OSError("DOCX renderer produced no output")
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _pdf_unicode_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    registered_name = "STSong-Light"
    if registered_name in pdfmetrics.getRegisteredFontNames():
        return registered_name
    pdfmetrics.registerFont(UnicodeCIDFont(registered_name))
    return registered_name


def write_pdf(payload: dict[str, Any], path: Path) -> None:
    """Write a derived Unicode PDF view from the canonical transcript payload."""

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    path = Path(path)
    if path.exists():
        raise FileExistsError(f"Transcript export already exists: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    font_name = _pdf_unicode_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "APMATitle",
        parent=styles["Title"],
        fontName=font_name,
        alignment=TA_CENTER,
        fontSize=18,
        leading=22,
    )
    body_style = ParagraphStyle(
        "APMABody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=9.5,
        leading=14,
        wordWrap="CJK",
    )
    meta_style = ParagraphStyle(
        "APMAMeta",
        parent=body_style,
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#39495A"),
    )
    document = SimpleDocTemplate(
        str(temporary),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="APMA Meeting Transcript",
        subject="Derived from APMA canonical FINAL JSON",
    )
    story = [
        Paragraph("Meeting transcript", title_style),
        Paragraph(
            "<b>Authority notice:</b> This PDF is a derived view. "
            "The corresponding FINAL.json is authoritative.",
            meta_style,
        ),
        Spacer(1, 4 * mm),
    ]
    metadata_rows = [
        [Paragraph(escape(label), meta_style), Paragraph(escape(value), meta_style)]
        for label, value in _metadata_items(payload)
    ]
    metadata = Table(metadata_rows, colWidths=[38 * mm, 118 * mm], repeatRows=0)
    metadata.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C0CA")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF3F8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([metadata, Spacer(1, 5 * mm)])
    for segment in payload.get("segments") or []:
        if not isinstance(segment, dict):
            continue
        text = str(segment.get("text", "")).strip()
        if not text:
            continue
        speaker = str(segment.get("speaker") or "Unlabelled speaker")
        story.append(
            KeepTogether(
                [
                    Paragraph(
                        f"<b>{escape(speaker)}</b> | {escape(_segment_timing(segment))}",
                        meta_style,
                    ),
                    Paragraph(escape(text).replace("\n", "<br/>"), body_style),
                    Spacer(1, 3 * mm),
                ]
            )
        )
    try:
        document.build(story)
        if not temporary.is_file() or not temporary.read_bytes().startswith(b"%PDF-"):
            raise OSError("PDF renderer produced invalid output")
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def write_transcript_exports(payload: dict[str, Any], outputs_dir: Path) -> dict[str, str]:
    if not isinstance(payload, dict):
        raise TypeError("Transcript payload must be a dictionary")

    outputs_dir = Path(outputs_dir)
    outputs_dir.mkdir(parents=True, exist_ok=True)
    paths = {key: outputs_dir / filename for key, filename in EXPORT_FILENAMES.items()}
    _write_text(paths["full_transcript_html"], render_html(payload))
    _write_text(paths["full_transcript_srt"], render_srt(payload))
    _write_text(paths["full_transcript_vtt"], render_vtt(payload))
    return {key: str(path) for key, path in paths.items()}
